import hashlib
import io
import tempfile
import unittest
from pathlib import Path
from copy import deepcopy
from PIL import Image,ImageChops,ImageStat
from pptx import Presentation
from converter import convert,read_package,inspect,xml,serialized,write_package,render,NS,ConversionError
from tests.fixtures import make_deck

class ConverterTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory();cls.base=Path(cls.tmp.name)
        cls.source=make_deck(cls.base/'original.pptx')
        cls.original=read_package(cls.source)
        cls.edit=cls.base/'editable.pptx';cls.visual=cls.base/'visual.pptx'
        cls.hash=hashlib.sha256(cls.source.read_bytes()).digest()
        cls.er=convert(cls.source,cls.edit)
        cls.images=render(cls.source,cls.base/'before',width=1200)
        cls.vr=convert(cls.source,cls.visual,'visual',images=cls.images)
        cls.after=render(cls.edit,cls.base/'after',width=1600)
        cls.visual_after=render(cls.visual,cls.base/'visual-after',width=1600)

    @classmethod
    def tearDownClass(cls):cls.tmp.cleanup()

    def test_geometry_all_parts_and_nested_groups(self):
        edited=read_package(self.edit)
        for name,data in self.original.items():
            if not name.endswith('.xml') or not name.startswith(('ppt/slides/','ppt/slideMasters/','ppt/slideLayouts/')):continue
            before=xml(data);after=xml(edited[name])
            for tag,attrs in [('off',['x','y']),('ext',['cx','cy']),('chOff',['x','y']),('chExt',['cx','cy']),('gridCol',['w']),('tr',['h'])]:
                bs=before.xpath(f'//*[local-name()="{tag}"]');as_=after.xpath(f'//*[local-name()="{tag}"]')
                self.assertEqual(len(bs),len(as_))
                for b,a in zip(bs,as_):
                    for attr in attrs:
                        if b.get(attr) is not None:
                            value=int(b.get(attr));expected=(value*4+1)//3 if attr in ('x','cx','w') else value
                            self.assertEqual(int(a.get(attr)),expected,(name,tag,attr))
        # Actual display coordinates use parent ext/chExt ratios, unchanged by conversion.
        original=Presentation(self.source);wide=Presentation(self.edit)
        g=original.slides[0].shapes[3];wg=wide.slides[0].shapes[3]
        def global_child_x(group):
            xf=group._element.find('p:grpSpPr/a:xfrm',NS)
            off=int(xf.find('a:off',NS).get('x'));ext=int(xf.find('a:ext',NS).get('cx'))
            cho=int(xf.find('a:chOff',NS).get('x'));che=int(xf.find('a:chExt',NS).get('cx'))
            return off+(group.shapes[0].left-cho)*ext/che
        self.assertAlmostEqual(global_child_x(wg),global_child_x(g)*4/3,delta=2)

    def test_text_and_preserved_parts(self):
        edited=read_package(self.edit)
        for name,data in self.original.items():
            if name.endswith('.rels') or name.startswith(('ppt/notes','ppt/media/','ppt/charts/','ppt/embeddings/')):
                self.assertEqual(data,edited[name],name)
            if name.startswith('ppt/slides/') and name.endswith('.xml'):
                b,a=xml(data),xml(edited[name])
                self.assertEqual(b.xpath('//a:t/text()',namespaces=NS),a.xpath('//a:t/text()',namespaces=NS))
                self.assertEqual(len(b.findall('.//a:br',NS)),len(a.findall('.//a:br',NS)))
                self.assertEqual(b.xpath('//@sz'),a.xpath('//@sz'))
                for tag in ['timing','transition']:
                    be=b.find('p:'+tag,NS);ae=a.find('p:'+tag,NS)
                    if be is not None:self.assertEqual(serialized(be),serialized(ae))
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).digest(),self.hash)

    def test_visual_pixels_full_bleed_and_height(self):
        parts=read_package(self.visual);info=inspect(parts)
        self.assertEqual(info['height'],inspect(self.original)['height'])
        self.assertEqual(len(self.images),3)
        for index,impath in enumerate(self.images):
            with Image.open(impath) as before:
                image=Image.open(io.BytesIO(parts[f'ppt/media/stretch-{index+1}.png']))
                self.assertEqual(image.height,before.height)
                self.assertEqual(image.width,round(before.width*4/3))
                expected=before.resize(image.size,Image.Resampling.LANCZOS)
                self.assertIsNone(ImageChops.difference(expected,image).getbbox())
            root=xml(parts[info['slides'][index]])
            self.assertIsNone(root.find('p:timing',NS))
            self.assertEqual(len(root.findall('p:cSld/p:spTree/p:pic',NS)),1)
            ext=root.find('.//p:pic/p:spPr/a:xfrm/a:ext',NS)
            self.assertEqual(int(ext.get('cx')),info['width']);self.assertEqual(int(ext.get('cy')),info['height'])
            # Rendered exported file must actually show the embedded stretched image.
            actual=Image.open(self.visual_after[index]).convert('RGB').resize(image.size)
            diff=ImageStat.Stat(ImageChops.difference(actual,image.convert('RGB')))
            self.assertLess(sum(diff.mean)/3,4.0)
        for n,b in self.original.items():
            if n.startswith('ppt/notes'):self.assertEqual(b,parts[n])

    def test_renderer_line_breaks_and_vertical_extent(self):
        import fitz
        def lines(path):
            with fitz.open(path) as doc:
                return [[(line['bbox'],''.join(s['text'] for s in line['spans'])) for block in page.get_text('dict')['blocks'] if 'lines' in block for line in block['lines']] for page in doc]
        before=lines(self.base/'before/render.pdf');after=lines(self.base/'after/render.pdf')
        # Explicit breaks remain lines; height/y of their rendered glyphs is unchanged.
        for text in ['First line stays here.','Second line stays here.','Third line stays here.']:
            b=next(box for box,t in before[0] if t==text);a=next(box for box,t in after[0] if t==text)
            self.assertAlmostEqual(b[1],a[1],delta=1);self.assertAlmostEqual(b[3]-b[1],a[3]-a[1],delta=1)
        for text in ['Bottom edge stays visible','No vertical scaling']:
            self.assertTrue(any(t==text for _,t in after[1]),text)
        # Automatic wrapping IS allowed to differ; explicitly prove/document the limitation.
        self.assertNotEqual([t for _,t in before[1]],[t for _,t in after[1]])
        # No missing words across multiline text, images and tables on representative slides.
        for i in (0,1):
            self.assertEqual(' '.join(' '.join(t for _,t in before[i]).split()).split(), ' '.join(' '.join(t for _,t in after[i]).split()).split())
        for page in after:
            for box,text in page:
                self.assertGreaterEqual(box[0],0,text);self.assertLessEqual(box[2],960,text)
                self.assertGreaterEqual(box[1],0,text);self.assertLessEqual(box[3],540,text)

    def test_already_wide_is_byte_identical(self):
        dest=self.base/'twice.pptx';report=convert(self.edit,dest,'visual')
        self.assertTrue(report['already_wide']);self.assertEqual(self.edit.read_bytes(),dest.read_bytes())

    def test_other_ratio_and_original_protection(self):
        with self.assertRaises(ConversionError):convert(self.source,self.source)
        parts=deepcopy(self.original);root=xml(parts['ppt/presentation.xml']);root.find('p:sldSz',NS).set('cx','9144000');root.find('p:sldSz',NS).set('cy','9144000');parts['ppt/presentation.xml']=serialized(root)
        square=self.base/'square.pptx';write_package(parts,square)
        with self.assertRaisesRegex(ConversionError,'Unsupported aspect'):convert(square,self.base/'bad.pptx')
        self.assertFalse((self.base/'bad.pptx').exists())

if __name__=='__main__':unittest.main()
