"""Representative 4:3 deck, generated locally without external assets."""
from pathlib import Path
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
from lxml import etree as E
from PIL import Image, ImageDraw
from converter import NS,read_package,write_package,xml,serialized

def make_deck(path):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    prs=Presentation();prs.slide_width=Inches(10);prs.slide_height=Inches(7.5)
    s=prs.slides.add_slide(prs.slide_layouts[6])
    bg=s.background.fill;bg.solid();bg.fore_color.rgb=RGBColor(247,248,235)
    def text(slide,x,y,w,h,content,size=24):
        shape=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.text_frame.word_wrap=True
        shape.text=content
        for p in shape.text_frame.paragraphs:
            p.font.size=Pt(size);p.font.name='DejaVu Sans';p.font.color.rgb=RGBColor(30,65,55)
        return shape
    title=text(s,.4,.25,8.8,.65,'SCIENCE · A wider classroom',28)
    title.text_frame.paragraphs[0].runs[0].hyperlink.address='https://example.com/lesson'
    text(s,.4,1.25,4.1,1.5,'First line stays here.\nSecond line stays here.\nThird line stays here.',20)
    image=Image.new('RGB',(300,180),'#c9dbb4');d=ImageDraw.Draw(image)
    d.rectangle((0,0,299,179),outline='#285e50',width=8);d.ellipse((90,30,210,150),fill='#edac51');d.line((0,90,300,90),fill='#285e50',width=3)
    img=path.parent/'fixture.png';image.save(img)
    s.shapes.add_picture(str(img),Inches(5.2), Inches(1.25), Inches(4.2), Inches(2.0))
    group=s.shapes.add_group_shape()
    group.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.6),Inches(3.4),Inches(1.6),Inches(.8))
    child=group.shapes.add_group_shape()
    child.shapes.add_shape(MSO_SHAPE.OVAL,Inches(2.5),Inches(3.4),Inches(1.2),Inches(.8))
    cell=s.shapes.add_table(3,3,Inches(.6),Inches(4.6),Inches(8.8),Inches(1.7))
    for row,values in zip(cell.table.rows,[['Topic','Before','After'],['Canvas','4:3','16:9'],['Height','Unchanged','Unchanged']]):
        for c,value in zip(row.cells,values):
            c.text=value
            for p in c.text_frame.paragraphs:p.font.size=Pt(18)
    s.notes_slide.notes_text_frame.text='Teacher note: original speaker notes must survive.'
    # A master-owned rectangle appears on all slides without a local duplicate.
    master_shape=deepcopy(group.shapes[0]._element)
    master_shape.find('p:nvSpPr/p:cNvPr',NS).set('id','999')
    xf=master_shape.find('p:spPr/a:xfrm',NS)
    xf.find('a:off',NS).set('x',str(Inches(8.9)));xf.find('a:off',NS).set('y',str(Inches(6.8)))
    xf.find('a:ext',NS).set('cx',str(Inches(.6)));xf.find('a:ext',NS).set('cy',str(Inches(.3)))
    prs.slide_master.shapes._spTree.append(master_shape)
    s2=prs.slides.add_slide(prs.slide_layouts[6])
    text(s2,.5,.4,9,.7,'Text wrapping & edge checks',28)
    text(s2,.5,1.5,4.2,2,'Automatic wrapping changes when a wider text box has room for more words. Review these line breaks.',24)
    text(s2,5.2,1.5,4.2,2,'Explicit first line\nExplicit second line\nExplicit third line',24)
    text(s2,.5,5.7,9,1,'Bottom edge stays visible\nNo vertical scaling',24)
    s3=prs.slides.add_slide(prs.slide_layouts[0]);s3.shapes.title.text='Inherited title & editable chart'
    s3.placeholders[1].text='Layout and master geometry are each scaled once.'
    chart=CategoryChartData();chart.categories=['A','B','C'];chart.add_series('Students',[8,12,10])
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1),Inches(3),Inches(8),Inches(3.6),chart)
    prs.save(path)
    parts=read_package(path)
    root=xml(parts['ppt/slides/slide1.xml'])
    transition=E.SubElement(root,'{'+NS['p']+'}transition',spd='slow')
    E.SubElement(transition,'{'+NS['p']+'}fade')
    # Minimal valid timing tree with a stable shape target.
    timing=E.SubElement(root,'{'+NS['p']+'}timing')
    tn=E.SubElement(timing,'{'+NS['p']+'}tnLst');par=E.SubElement(tn,'{'+NS['p']+'}par')
    ct=E.SubElement(par,'{'+NS['p']+'}cTn',id='1',dur='indefinite',restart='never',nodeType='tmRoot')
    children=E.SubElement(ct,'{'+NS['p']+'}childTnLst')
    effect=E.SubElement(children,'{'+NS['p']+'}animEffect',transition='in',filter='fade')
    behavior=E.SubElement(effect,'{'+NS['p']+'}cBhvr')
    E.SubElement(behavior,'{'+NS['p']+'}cTn',id='2',dur='500',fill='hold')
    target=E.SubElement(behavior,'{'+NS['p']+'}tgtEl')
    E.SubElement(target,'{'+NS['p']+'}spTgt',spid='2')
    parts['ppt/slides/slide1.xml']=serialized(root);write_package(parts,path)
    return path

if __name__=='__main__':
    make_deck(Path('artifacts/representative-4x3.pptx'))
